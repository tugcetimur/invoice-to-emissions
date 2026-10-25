import pytest
from openpyxl import load_workbook

from src.report import results_table, review_table, summary_table, write_workbook

ELECTRICITY = {
    "invoice_type": "electricity", "fuel_type": None, "supplier": "Northwind Energy",
    "period_start": "2025-04-01", "period_end": "2025-04-30",
    "quantity": 404.0, "unit": "kWh", "total_amount": 115.3, "currency": "EUR",
}


def processed_result(file="a.pdf"):
    return {
        "file": file,
        "status": "processed",
        "reasons": [],
        "suggestions": {},
        "fields": dict(ELECTRICITY),
        "candidates": {name: (value, value) for name, value in ELECTRICITY.items()},
        "emission": {
            "quantity": 404.0, "unit": "kWh", "factor": 0.177, "factor_unit": "kWh",
            "scope": 2, "kg_co2e": 71.508, "t_co2e": 0.071508,
            "source": "UK DESNZ", "year": 2025,
        },
    }


def review_result(file="b.pdf"):
    candidates = {name: (value, value) for name, value in ELECTRICITY.items()}
    candidates["quantity"] = (4100.0, 520.0)
    return {
        "file": file,
        "status": "needs_review",
        "reasons": ["quantity: conflict (regex=4100.0, llm=520.0); only 520.0 passes the checks"],
        "suggestions": {"quantity": 520.0},
        "fields": {**ELECTRICITY, "quantity": None},
        "candidates": candidates,
        "emission": None,
    }


def test_results_table_has_only_processed_invoices():
    table = results_table([processed_result("a.pdf"), review_result("b.pdf")])
    assert list(table["file"]) == ["a.pdf"]
    assert table.loc[0, "kg_co2e"] == pytest.approx(71.508)
    assert table.loc[0, "scope"] == 2


def test_review_table_keeps_reasons_candidates_and_suggestions():
    table = review_table([processed_result("a.pdf"), review_result("b.pdf")])
    row = table.iloc[0]
    assert row["file"] == "b.pdf"
    assert "conflict" in row["reasons"]
    assert row["suggestion"] == "quantity = 520.0"
    assert row["quantity_regex"] == 4100.0
    assert row["quantity_llm"] == 520.0


def test_summary_counts_and_totals():
    results = [processed_result("a.pdf"), processed_result("c.pdf"), review_result("b.pdf")]
    table = summary_table(results)
    summary = dict(zip(table["metric"], table["value"]))
    assert summary["Invoices"] == 3
    assert summary["Processed automatically"] == 2
    assert summary["Sent to review"] == 1
    assert summary["Scope 1 (tCO2e)"] == pytest.approx(0.0)
    assert summary["Scope 2 (tCO2e)"] == pytest.approx(0.143)
    assert summary["Total (tCO2e)"] == pytest.approx(0.143)


def test_empty_results_still_produce_tables():
    assert len(results_table([])) == 0
    assert "kg_co2e" in results_table([]).columns
    assert len(review_table([])) == 0
    table = summary_table([])
    assert dict(zip(table["metric"], table["value"]))["Invoices"] == 0


def test_workbook_has_three_sheets_with_the_right_rows(tmp_path):
    path = write_workbook([processed_result("a.pdf"), review_result("b.pdf")], tmp_path / "report.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Results", "Review queue", "Summary"]
    assert [cell.value for cell in workbook["Results"][1]][:3] == ["file", "invoice_type", "fuel_type"]
    assert workbook["Results"].max_row == 2
    assert workbook["Review queue"].max_row == 2
    assert workbook["Review queue"]["A2"].value == "b.pdf"
    summary = dict(list(workbook["Summary"].iter_rows(values_only=True))[1:])
    assert summary["Invoices"] == 2
    assert summary["Processed automatically"] == 1
    assert summary["Sent to review"] == 1


def test_workbook_creates_missing_folders(tmp_path):
    path = write_workbook([processed_result()], tmp_path / "out" / "deep" / "report.xlsx")
    assert path.exists()
